from docx import Document
from pathlib import Path
import json

paths = [
    Path(r"C:\Users\блоош\Downloads\Telegram Desktop\РПМ Л9 (2).docx"),
    Path(r"C:\Users\блоош\Downloads\Telegram Desktop\РПМ Л22.docx"),
]

for path in paths:
    doc = Document(path)
    data = {
        "file": path.name,
        "sections": [
            {
                "width": s.page_width,
                "height": s.page_height,
                "top": s.top_margin,
                "bottom": s.bottom_margin,
                "left": s.left_margin,
                "right": s.right_margin,
            }
            for s in doc.sections
        ],
        "paragraphs": [
            {"style": p.style.name, "text": p.text}
            for p in doc.paragraphs if p.text.strip()
        ],
        "tables": [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in doc.tables
        ],
    }
    print(json.dumps(data, ensure_ascii=True))

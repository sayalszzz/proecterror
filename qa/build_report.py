import json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

GOOD = Path(r"C:\Users\блоош\IdeaProjects\proect")
BAD = Path(r"C:\Users\блоош\IdeaProjects\proecterror")
OUT = BAD / "Отчет_группа_6_ресторан.docx"
GOOD_SCREEN = BAD / "qa" / "proect_gui_final.png"
BAD_SCREEN = BAD / "qa" / "proecterror_gui_final.png"
BENCHMARK = json.loads((BAD / "qa" / "benchmark-results.json").read_text(encoding="utf-8"))

def font(run, size=14, bold=False, name="Times New Roman", color=None):
    run.font.name = name
    props = run._element.get_or_add_rPr()
    fonts = props.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); props.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)
    run.font.size = Pt(size); run.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)

def field(paragraph, code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = code
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end]); font(run, size=10)

def body(doc, text, indent=True):
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if indent: p.paragraph_format.first_line_indent = Cm(1.25)
    font(p.add_run(text)); return p

def heading(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(6)
    font(p.add_run(text), size=14, bold=True); return p

def shade(cell, fill):
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, text in enumerate(headers):
        cell = t.rows[0].cells[i]; shade(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(text), size=10, bold=True, color=(255,255,255))
    for row_no, values in enumerate(rows):
        cells = t.add_row().cells
        for i, text in enumerate(values):
            if row_no % 2: shade(cells[i], "EAF2F8")
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            font(p.add_run(str(text)), size=9)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths): row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def code(doc, title, lines):
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
    font(p.add_run(title), size=11, bold=True)
    for line in lines:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(.6)
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1
        font(p.add_run(line.rstrip()), size=8, name="Consolas")

def caption(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(8)
    font(p.add_run(text), size=11)

doc = Document(); sec = doc.sections[0]
sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
sec.left_margin = Cm(2.5); sec.right_margin = Cm(1.5)
normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(14)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; field(footer, "PAGE")

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Cm(2.5)
font(p.add_run("ОТЧЁТ\nПО ПРОЕКТНОЙ РАБОТЕ"), size=18, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Cm(1.2)
font(p.add_run("Вариант (группа) 6\n«Управление меню и заказами ресторана»"), size=16, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Cm(1.5)
font(p.add_run("Две версии программного продукта:\nproect — исправленная; proecterror — с учебными ошибками"))
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_before = Cm(5)
font(p.add_run("Выполнил(а): ____________________\nПроверил(а): ____________________"))
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Cm(3)
font(p.add_run("2026")); doc.add_page_break()

heading(doc, "1. Цель и постановка задачи")
body(doc, "Цель работы — разработать Java-приложение для управления меню и заказами ресторана, подготовить две функционально сопоставимые версии и показать на их сравнении ошибки рефакторинга и оптимизации.")
body(doc, "В соответствии с заданием варианта 6 реализованы два проекта: proect с исправленной архитектурой и proecterror с полным учебным набором из 11 ошибок рефакторинга и 12 ошибок оптимизации. Обе версии используют графический интерфейс JavaFX и постоянное хранение данных в SQLite.")
heading(doc, "2. Функциональные требования")
table(doc, ["Функция", "proect", "proecterror"], [
    ("Просмотр меню", "Реализовано", "Реализовано"), ("Добавление блюда", "С проверкой данных", "Реализовано"),
    ("Поиск по названию", "Реализовано", "Реализовано"), ("Фильтр по категории", "Реализовано", "Реализовано"),
    ("Выбор нескольких блюд", "Множественный выбор", "Множественный выбор"), ("Создание заказа", "Реализовано", "Реализовано"),
    ("Удаление/очистка позиций", "Реализовано", "Удаление позиции"), ("История заказов", "SQLite, сохраняется", "SQLite, сохраняется"),
    ("Изменение статуса", "Выбранный заказ истории", "Выбранный заказ истории"), ("Расчёт выручки", "SQL-агрегация", "Перебор всех заказов"),
    ("Пагинация", "По 10 записей", "Отсутствует"),
], [5.4,5.1,5.1])

heading(doc, "3. Средства разработки")
body(doc, "Язык программирования — Java 17. Для пользовательского интерфейса применена JavaFX 17.0.12, для доступа к базе данных — JDBC-драйвер Xerial sqlite-jdbc 3.42.0.0. Сборка описана в Maven-файлах pom.xml; точкой входа служит класс Main.")
heading(doc, "4. Архитектура исправленного проекта proect")
table(doc, ["Класс", "Назначение"], [
    ("Dish", "Неизменяемая модель блюда; проверяет идентификатор, название, категорию и цену."),
    ("NutritionInfo", "Группирует вес и калорийность блюда."), ("Customer", "Группирует имя клиента и номер телефона."),
    ("Order", "Хранит позиции, клиента, статус, время и вычисленную стоимость."), ("OrderStatus", "Перечисляет допустимые статусы."),
    ("RestaurantDAO", "Создаёт схему SQLite и выполняет параметризованные SQL-запросы."),
    ("RestaurantController", "Содержит прикладные операции над блюдами и заказами."),
    ("RestaurantApp", "Формирует JavaFX-интерфейс; запускает обращения к SQLite в Task."), ("Main", "Запускает RestaurantApp."),
], [4.2,11.4])

heading(doc, "5. Хранение данных в SQLite")
body(doc, "В проекте proect используется файл restaurant.db, в proecterror — restaurant-error.db. При первом запуске создаются таблицы dishes, orders и order_items. Начальные блюда добавляются только в пустую базу, поэтому повторный запуск не удаляет историю заказов.")
table(doc, ["Таблица", "Основные данные"], [("dishes", "Блюда меню и характеристики."), ("orders", "Клиент, телефон, статус, время и итог."), ("order_items", "Связь заказа с выбранными блюдами.")], [4.2,11.4])
body(doc, "В исправленной версии SQL выполняется через PreparedStatement, создание заказа защищено транзакцией, а загрузка заказов и позиций выполняется одним JOIN-запросом. Меню и история загружаются страницами по 10 записей. Созданы индексы по категории, названию, order_id, статусу и времени заказа.")

heading(doc, "6. Отсутствие зависания интерфейса в proect")
body(doc, "Все операции, которые обращаются к SQLite, передаются в фоновый javafx.concurrent.Task. После завершения результат возвращается в setOnSucceeded, который безопасно обновляет интерфейс. Во время операции выводится сообщение «Работа с базой...». Длительный запрос поэтому не блокирует окно.")
good_app = (GOOD / "src" / "RestaurantApp.java").read_text(encoding="utf-8").splitlines()
start = next(i for i,line in enumerate(good_app) if "private <T> void runDb" in line)
code(doc, "Листинг 1 — запуск работы с БД в фоновом Task", good_app[start:start+16])

doc.add_page_break(); heading(doc, "7. Интерфейс исправленного проекта")
body(doc, "Окно proect содержит меню с поиском и SQL-фильтром, форму добавления блюда, данные клиента, текущий заказ, историю, изменение статуса и общую выручку. Меню и история имеют кнопки пагинации. Несколько строк выбираются с помощью Ctrl или Shift и добавляются одной кнопкой, а запись в SQLite происходит после нажатия «Оформить заказ».")
doc.add_picture(str(GOOD_SCREEN), width=Cm(16.2)); caption(doc, "Рисунок 1 — интерфейс исправленного проекта proect после повторного запуска")

doc.add_page_break(); heading(doc, "8. Интерфейс проекта с учебными ошибками")
body(doc, "В proecterror доступны добавление блюда, поиск, фильтр категорий, множественный выбор, оформление и история. Для смены статуса требуется выделить заказ в правом списке, выбрать статус и нажать «Изменить статус выбранного заказа». База при запуске не очищается.")
doc.add_picture(str(BAD_SCREEN), width=Cm(16.2)); caption(doc, "Рисунок 2 — интерфейс proecterror с фильтром и сменой статуса")

doc.add_page_break(); heading(doc, "9. Ошибки рефакторинга в proecterror")
table(doc, ["№", "Ошибка", "Исправление в proect"], [
    (1,"Публичные поля моделей нарушают инкапсуляцию.","private final и методы доступа."),
    (2,"Конструкторы Dish и Order имеют 8 и 6 параметров.","Питание и клиент вынесены в отдельные типы."),
    (3,"Создание Dish из ResultSet продублировано.","Применяется единый mapDish()."),
    (4,"getAllOrders() превышает 40 строк и смешивает обязанности.","Загрузка разделена на короткие методы."),
    (5,"Статусы и число 100 заданы магическими значениями.","OrderStatus и PERCENT_BASE."),
    (6,"Feature Envy в getDailyStats().","Расчёт инкапсулирован в Order и DAO."),
    (7,"customerName/customerPhone и weight/calories образуют Data Clumps.","Customer и NutritionInfo."),
    (8,"Входные данные не проверяются.","Валидация в моделях и интерфейсе."),
    (9,"SQLException только печатается в консоль.","Ошибка передаётся интерфейсу с причиной."),
    (10,"Связанные параметры не объединены.","Применены объекты-значения."),
    (11,"Комментарий «Настройка таблиц» ничего не объясняет.","Код схемы разделён и документирован по назначению."),
], [1,7.1,8])
heading(doc, "10. Ошибки оптимизации в proecterror")
table(doc, ["№", "Ошибка", "Исправление в proect"], [
    (1,"N+1 при загрузке позиций заказов.","Один JOIN-запрос."),
    (2,"Нет PreparedStatement, возможна SQL-инъекция.","Параметризованные запросы."),
    (3,"Все записи загружаются без пагинации.","LIMIT/OFFSET и страницы по 10 записей."),
    (4,"Нет индексов по связям, статусу и времени.","Добавлены пять предметных индексов."),
    (5,"Категория фильтруется линейно.","Фильтрация выполняется индексированным SQL."),
    (6,"toLowerCase() вычисляется многократно.","Нормализация вынесена из повторяющегося кода."),
    (7,"Заказ и позиции записываются без транзакции.","commit/rollback охватывает обе операции."),
    (8,"JDBC-ресурсы могут утекать.","try-with-resources для Connection, Statement и ResultSet."),
    (9,"Итог полностью пересчитывается при изменении статуса.","Неизменяемый Order сохраняет рассчитанный total."),
    (10,"Повторные запросы выполняются без кэша.","Кэш страниц меню с инвалидированием."),
    (11,"SQL собирается конкатенацией.","StringBuilder и параметры."),
    (12,"База вызывается из JavaFX-потока.","Все обращения выполняются в Task."),
], [1,7.1,8])

doc.add_page_break(); heading(doc, "11. Измерение улучшений")
body(doc, "Для сравнения создан воспроизводимый тест на SQLite с 5000 блюдами, 1000 заказами и 5000 позициями. В таблице приведена медиана семи запусков на текущем компьютере. Результаты зависят от оборудования, но показывают влияние выбранных алгоритмов.")
table(doc, ["Сценарий", "До, мс", "После, мс", "Ускорение"], [
    (item["name"], f'{item["before"]:.3f}', f'{item["after"]:.3f}', f'{item["speedup"]:.1f}x')
    for item in BENCHMARK["scenarios"]
], [7.2,2.8,2.8,3.0])
body(doc, "Наибольший эффект дал кэш страницы меню. JOIN устранил сотни отдельных запросов, а индексированный SQL исключил загрузку 5000 строк ради одной категории. Скрипт qa/benchmark.py позволяет повторить измерение.")

heading(doc, "12. Сравнение проектов")
body(doc, "Проекты имеют одну предметную область и сопоставимый интерфейс, но различаются внутренним устройством. proect демонстрирует инкапсуляцию, объекты-значения, enum, транзакции, PreparedStatement, JOIN, индексы, пагинацию, кэш и фоновые Task. proecterror остаётся работоспособным специально для демонстрации полного набора проблем из методички.")
heading(doc, "13. Проверка работоспособности")
table(doc, ["Проверка", "Результат"], [
    ("Компиляция proect", "Успешно, код javac 0."), ("Компиляция proecterror", "Успешно, код javac 0."),
    ("Запуск proect", "Окно открыто; история и выручка загружены из restaurant.db."),
    ("Запуск proecterror", "Окно открыто; история сохранена в restaurant-error.db."),
    ("Повторный запуск", "Начальные данные не дублируются, таблицы не очищаются."),
    ("Фильтр и поиск", "Присутствуют в обеих версиях."),
    ("Смена статуса в proecterror", "Выбирается заказ истории, статус сохраняется через DAO."),
    ("Пагинация proect", "Кнопки страниц присутствуют для меню и истории."),
    ("Тест производительности", "Три сценария, медиана семи запусков."),
], [8.2,7.9])
body(doc, "При JDK 24 SQLite может выводить предупреждение native access. Оно не мешает запуску. Для его скрытия указывается VM option --enable-native-access=ALL-UNNAMED.")
heading(doc, "14. Порядок запуска")
body(doc, "В IntelliJ IDEA нужно открыть папку проекта, дождаться загрузки Maven-зависимостей, выбрать src/Main.java и выполнить Run 'Main.main()'. Сначала можно показать proect, затем proecterror. SQLite-файлы создаются рядом с проектами автоматически.")
heading(doc, "15. Теоретическая основа")
body(doc, "Для аргументации на защите использованы следующие материалы: Martin Fowler, Refactoring (https://refactoring.com); Oracle, Using Prepared Statements (https://docs.oracle.com/javase/tutorial/jdbc/basics/prepared.html); SQLite, Query Planning and Indexes (https://www.sqlite.org/queryplanner.html); OpenJFX, Concurrency in JavaFX (https://openjfx.io/javadoc/17/javafx.graphics/javafx/concurrent/Task.html); OWASP, SQL Injection Prevention Cheat Sheet (https://cheatsheetseries.owasp.org/cheatsheets/SQL_Injection_Prevention_Cheat_Sheet.html).")
heading(doc, "Заключение")
body(doc, "Задание варианта 6 выполнено: созданы две JavaFX-версии ресторана с SQLite и одинаковой предметной логикой. В proecterror воспроизведены и отмечены 11 ошибок рефакторинга и 12 ошибок оптимизации из методички. В proect они устранены с помощью выделения классов и методов, инкапсуляции, объектов-значений, PreparedStatement, транзакций, JOIN, индексов, пагинации, кэширования и фоновых Task. Корректность запуска подтверждена, а влияние оптимизаций измерено воспроизводимым тестом.")

doc.save(OUT); print(OUT)
